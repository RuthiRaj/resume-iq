"""
Document Extraction Service for ResumeIQ

Extracts raw text safely from uploaded document binaries (PDF, DOCX, TXT).
Enforces file size limits, extension allowlists, sanitization, and corrupt file handling.
"""

import io
import re
from typing import Tuple
from fastapi import HTTPException, status

MAX_DOCUMENT_SIZE_BYTES = 15 * 1024 * 1024  # 15 MB
MAX_EXTRACTED_TEXT_CHARS = 150_000  # ~30,000 words max

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt"}
PROHIBITED_EXTENSIONS = {
    ".exe", ".sh", ".bat", ".cmd", ".js", ".html", ".htm",
    ".svg", ".php", ".vbs", ".py", ".rb", ".dll", ".so",
}


def sanitize_filename(filename: str) -> str:
    """Removes path traversal characters and normalizes filename."""
    if not filename:
        return "unnamed_document"
    clean = filename.replace("\\", "/").split("/")[-1]
    clean = re.sub(r"[^a-zA-Z0-9._-]", "_", clean)
    return clean[:100] or "unnamed_document"


def validate_document_upload(filename: str, content_length: int) -> str:
    """Validates document extension, name, and size boundaries."""
    if content_length <= 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded file is empty (0 bytes).",
        )

    if content_length > MAX_DOCUMENT_SIZE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_CONTENT_TOO_LARGE,
            detail=f"File exceeds maximum allowed size of {MAX_DOCUMENT_SIZE_BYTES // (1024 * 1024)}MB.",
        )

    clean_name = sanitize_filename(filename)
    lower_name = clean_name.lower()

    if any(lower_name.endswith(ext) for ext in PROHIBITED_EXTENSIONS):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Executable or script files are strictly prohibited.",
        )

    ext = "." + lower_name.split(".")[-1] if "." in lower_name else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '{ext}'. Supported formats: PDF, DOCX, TXT.",
        )

    return clean_name


def extract_text_from_pdf(content: bytes) -> str:
    """Extracts text from PDF binary content using pypdf."""
    try:
        from pypdf import PdfReader

        stream = io.BytesIO(content)
        reader = PdfReader(stream)

        if reader.is_encrypted:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Encrypted or password-protected PDF files cannot be parsed. Please upload an unlocked PDF.",
            )

        text_parts = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(page_text.strip())

        extracted = "\n\n".join(text_parts).strip()
        if not extracted:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail="No readable text could be extracted from the PDF. Scanned images without selectable text are not supported.",
            )

        return extracted
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Failed to parse PDF document. The file may be corrupt or malformed.",
        ) from exc


def extract_text_from_docx(content: bytes) -> str:
    """Extracts text from DOCX binary content using python-docx."""
    try:
        import docx

        stream = io.BytesIO(content)
        doc = docx.Document(stream)

        text_parts = []
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                text_parts.append(p.text.strip())

        # Also extract table cells
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    text_parts.append(" | ".join(row_texts))

        extracted = "\n".join(text_parts).strip()
        if not extracted:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail="No readable text found in the DOCX document.",
            )

        return extracted
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Failed to parse DOCX document. The file may be corrupt or malformed.",
        ) from exc


def extract_text_from_txt(content: bytes) -> str:
    """Extracts text from raw TXT bytes."""
    for encoding in ["utf-8", "latin-1", "cp1252"]:
        try:
            text = content.decode(encoding).strip()
            if text:
                return text
        except UnicodeDecodeError:
            continue

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="Failed to decode text document. Unsupported character encoding.",
    )


def extract_document_text(filename: str, content: bytes) -> str:
    """
    Unified extraction entrypoint.
    Validates file boundaries, detects format, and extracts normalized text.
    """
    clean_name = validate_document_upload(filename, len(content))
    lower_name = clean_name.lower()

    if lower_name.endswith(".pdf"):
        raw_text = extract_text_from_pdf(content)
    elif lower_name.endswith(".docx") or lower_name.endswith(".doc"):
        raw_text = extract_text_from_docx(content)
    elif lower_name.endswith(".txt"):
        raw_text = extract_text_from_txt(content)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported document format.",
        )

    # Normalize excessive consecutive whitespaces and limit maximum length
    normalized = re.sub(r"\n{3,}", "\n\n", raw_text)
    normalized = re.sub(r"[ \t]{2,}", " ", normalized)

    if len(normalized) > MAX_EXTRACTED_TEXT_CHARS:
        normalized = normalized[:MAX_EXTRACTED_TEXT_CHARS]

    return normalized.strip()
